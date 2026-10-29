#!/usr/bin/env python3
"""
Собирает reference.docx для pandoc под типичные требования российских вузов
(ГОСТ 7.32-2017 / ГОСТ 2.105 / методички): Times New Roman 14, интервал 1.5,
абзацный отступ 1.25 см, выравнивание по ширине, А4.

Запуск:
    pandoc -o ref_default.docx --print-default-data-file reference.docx
    python3 make_reference.py ref_default.docx reference.docx

Все числа, которые обычно различаются между вузами, вынесены в блок НАСТРОЙКИ.
"""
import sys
from copy import deepcopy

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm

# ---------------------------- НАСТРОЙКИ ----------------------------
FONT = "Times New Roman"
CODE_FONT = "Courier New"
SIZE = 28                # размер в half-points: 28 = 14 pt
CODE_SIZE = 24           # 12 pt для кода
FOOTNOTE_SIZE = 20       # 10 pt
TABLE_SIZE = 24          # шрифт в таблицах: 24 = 12 pt
TABLE_LINE = 240         # интервал в таблицах: одинарный
LINE = 360               # 360 = 1.5 интервал (240 = одинарный)
INDENT = 709             # абзацный отступ 1.25 см в twips
MARGIN_LEFT_MM = 30
MARGIN_RIGHT_MM = 15     # во многих вузах 10
MARGIN_TOP_MM = 20
MARGIN_BOTTOM_MM = 20
H1_CAPS = True           # заголовок 1 уровня ПРОПИСНЫМИ
H1_CENTER = True         # и по центру
PAGE_NUMBER = True       # номер страницы внизу по центру
# --------------------------------------------------------------------

W = nsdecls("w")


def rpr(size=SIZE, bold=False, italic=False, caps=False, font=FONT,
        color="000000", underline=False):
    parts = [
        f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:eastAsia="{font}" w:cs="{font}"/>'
    ]
    if bold:
        parts.append("<w:b/><w:bCs/>")
    if italic:
        parts.append("<w:i/><w:iCs/>")
    if caps:
        parts.append("<w:caps/>")
    parts.append(f'<w:color w:val="{color}"/>')
    parts.append(f'<w:sz w:val="{size}"/><w:szCs w:val="{size}"/>')
    if underline:
        parts.append('<w:u w:val="single"/>')
    return f"<w:rPr {W}>{''.join(parts)}</w:rPr>"


def ppr(jc="both", first=0, left=0, hanging=0, line=LINE, before=0, after=0,
        keep_next=False, keep_lines=False, outline=None):
    parts = []
    if keep_next:
        parts.append("<w:keepNext/>")
    if keep_lines:
        parts.append("<w:keepLines/>")
    parts.append("<w:widowControl/>")
    parts.append(
        f'<w:spacing w:before="{before}" w:after="{after}" '
        f'w:line="{line}" w:lineRule="auto"/>'
    )
    ind = f'<w:ind w:left="{left}" w:right="0"'
    if hanging:
        ind += f' w:hanging="{hanging}"'
    else:
        ind += f' w:firstLine="{first}"'
    ind += "/>"
    parts.append(ind)
    parts.append(f'<w:jc w:val="{jc}"/>')
    if outline is not None:
        parts.append(f'<w:outlineLvl w:val="{outline}"/>')
    return f"<w:pPr {W}>{''.join(parts)}</w:pPr>"


def get_style(styles_el, style_id):
    for st in styles_el.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == style_id:
            return st
    return None


def set_style(styles_el, style_id, p=None, r=None):
    st = get_style(styles_el, style_id)
    if st is None:
        return None
    for tag in ("w:pPr", "w:rPr"):
        for old in st.findall(qn(tag)):
            st.remove(old)
    # pPr/rPr должны идти после qFormat и до tblPr/tblStylePr
    anchor = st.find(qn("w:tblPr"))
    for xml in (p, r):
        if xml is None:
            continue
        el = parse_xml(xml)
        if anchor is not None:
            anchor.addprevious(el)
        else:
            st.append(el)
    return st


def main(src, dst):
    doc = Document(src)
    styles_el = doc.styles.element

    # ---- docDefaults: шрифт, язык, интервалы ----
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is not None:
        styles_el.remove(dd)
    dd = parse_xml(
        f"<w:docDefaults {W}>"
        f'<w:rPrDefault><w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" '
        f'w:eastAsia="{FONT}" w:cs="{FONT}"/>'
        f'<w:sz w:val="{SIZE}"/><w:szCs w:val="{SIZE}"/>'
        f'<w:lang w:val="ru-RU" w:eastAsia="ru-RU" w:bidi="ar-SA"/></w:rPr></w:rPrDefault>'
        f'<w:pPrDefault><w:pPr><w:spacing w:after="0" w:line="{LINE}" w:lineRule="auto"/>'
        f"</w:pPr></w:pPrDefault></w:docDefaults>"
    )
    styles_el.insert(0, dd)

    body = ppr(first=INDENT)
    plain = ppr(first=0, jc="both")
    center = ppr(jc="center", first=0)

    # ---- основной текст ----
    set_style(styles_el, "Normal", body, rpr())
    set_style(styles_el, "BodyText", body, rpr())
    set_style(styles_el, "FirstParagraph", body, rpr())
    # Compact: ячейки таблиц и «плотные» списки
    set_style(styles_el, "Compact",
              ppr(first=0, jc="left", line=TABLE_LINE),
              rpr(size=TABLE_SIZE))
    set_style(styles_el, "BlockText", ppr(first=0, left=INDENT, jc="both"), rpr())
    set_style(styles_el, "Bibliography", ppr(first=INDENT), rpr())
    set_style(styles_el, "Definition", ppr(first=0, jc="both"), rpr())
    set_style(styles_el, "DefinitionTerm", ppr(first=0, jc="left", keep_next=True),
              rpr(bold=True))

    # ---- титульные элементы ----
    set_style(styles_el, "Title", ppr(jc="center", first=0, after=240, keep_next=True),
              rpr(bold=True, caps=True))
    set_style(styles_el, "Subtitle", ppr(jc="center", first=0, after=240, keep_next=True),
              rpr())
    set_style(styles_el, "Author", center, rpr())
    set_style(styles_el, "Date", center, rpr())
    set_style(styles_el, "AbstractTitle", ppr(jc="center", first=0, keep_next=True),
              rpr(bold=True, caps=True))
    set_style(styles_el, "Abstract", body, rpr())

    # ---- заголовки ----
    h1_jc = "center" if H1_CENTER else "left"
    h1_first = 0 if H1_CENTER else INDENT
    set_style(styles_el, "Heading1",
              ppr(jc=h1_jc, first=h1_first, before=240, after=240,
                  keep_next=True, keep_lines=True, outline=0),
              rpr(bold=True, caps=H1_CAPS))
    for lvl in range(2, 10):
        set_style(styles_el, f"Heading{lvl}",
                  ppr(jc="left", first=INDENT, before=240, after=120,
                      keep_next=True, keep_lines=True, outline=lvl - 1),
                  rpr(bold=True))
    set_style(styles_el, "TOCHeading",
              ppr(jc="center", first=0, keep_next=True),
              rpr(bold=True, caps=True))

    # ---- подписи, рисунки, таблицы ----
    set_style(styles_el, "Caption", center, rpr())
    # Рисунок: подпись ПОД рисунком, по центру
    set_style(styles_el, "ImageCaption",
              ppr(jc="center", first=0, before=120, after=240, line=240), rpr())
    # Таблица: название НАД таблицей, слева без отступа
    set_style(styles_el, "TableCaption",
              ppr(jc="left", first=0, before=120, after=120, line=240, keep_next=True),
              rpr())
    set_style(styles_el, "Figure", ppr(jc="center", first=0, line=240, keep_next=True), rpr())
    set_style(styles_el, "CaptionedFigure",
              ppr(jc="center", first=0, line=240, keep_next=True), rpr())

    # Table: сетка по всем границам
    tbl = get_style(styles_el, "Table")
    if tbl is not None:
        for old in tbl.findall(qn("w:tblPr")):
            tbl.remove(old)
        for old in tbl.findall(qn("w:tblStylePr")):
            tbl.remove(old)
        borders = "".join(
            f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            for side in ("top", "left", "bottom", "right", "insideH", "insideV")
        )
        tbl.append(parse_xml(
            f"<w:tblPr {W}><w:tblInd w:w=\"0\" w:type=\"dxa\"/>"
            f"<w:tblBorders>{borders}</w:tblBorders>"
            f'<w:tblCellMar><w:top w:w="0" w:type="dxa"/><w:left w:w="108" w:type="dxa"/>'
            f'<w:bottom w:w="0" w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar>'
            f"</w:tblPr>"
        ))
        tbl.append(parse_xml(
            f'<w:tblStylePr {W} w:type="firstRow"><w:rPr><w:b/><w:bCs/></w:rPr></w:tblStylePr>'
        ))

    # ---- сноски ----
    set_style(styles_el, "FootnoteText", ppr(first=INDENT, line=240),
              rpr(size=FOOTNOTE_SIZE))

    # ---- символьные стили ----
    set_style(styles_el, "BodyTextChar", None, rpr())
    set_style(styles_el, "VerbatimChar", None, rpr(size=CODE_SIZE, font=CODE_FONT))
    set_style(styles_el, "SectionNumber", None, rpr(bold=True))
    set_style(styles_el, "Hyperlink", None, rpr(underline=True))
    fr = get_style(styles_el, "FootnoteReference")
    if fr is not None:
        for old in fr.findall(qn("w:rPr")):
            fr.remove(old)
        fr.append(parse_xml(f'<w:rPr {W}><w:vertAlign w:val="superscript"/></w:rPr>'))

    # ---- SourceCode: блоки кода (pandoc подхватывает этот стиль) ----
    if get_style(styles_el, "SourceCode") is None:
        styles_el.append(parse_xml(
            f'<w:style {W} w:type="paragraph" w:customStyle="1" w:styleId="SourceCode">'
            f'<w:name w:val="Source Code"/><w:basedOn w:val="Normal"/>'
            f'<w:link w:val="VerbatimChar"/><w:qFormat/>'
            f'{ppr(jc="left", first=0, line=240, before=120, after=120)}'
            f'{rpr(size=CODE_SIZE, font=CODE_FONT)}</w:style>'
        ))

    # ---- страница: А4 и поля ----
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.left_margin = Mm(MARGIN_LEFT_MM)
    sec.right_margin = Mm(MARGIN_RIGHT_MM)
    sec.top_margin = Mm(MARGIN_TOP_MM)
    sec.bottom_margin = Mm(MARGIN_BOTTOM_MM)
    sec.header_distance = Mm(12.5)
    sec.footer_distance = Mm(10)

    # ---- номер страницы внизу по центру ----
    if PAGE_NUMBER:
        footer = sec.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        p.alignment = 1  # center
        p.paragraph_format.first_line_indent = 0
        for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
            run = p.add_run()
            run.font.name = FONT
            run.font.size = Mm(0) if False else None
            if kind:
                fc = parse_xml(f'<w:fldChar {W} w:fldCharType="{kind}"/>')
                run._r.append(fc)
            else:
                it = parse_xml(f"<w:instrText {W} xml:space=\"preserve\">{text}</w:instrText>")
                run._r.append(it)

    doc.save(dst)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("usage: make_reference.py ref_default.docx reference.docx")
    main(sys.argv[1], sys.argv[2])
